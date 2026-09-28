"""Build Miku's Stationery - a made-up kawaii stationery and book shop.

A test catalogue for the agent: a price table with a picture on most rows (the
photos the agent can send), a few rows without one, and shop terms written the
way a small shop writes them. Run from the backend virtualenv:

    python docs/test-stores/make_mikus_catalogue.py
"""

from __future__ import annotations

import io
import math
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).parent
OUT = HERE / "Mikus_Stationery_Catalogue.docx"

PINK, MINT, LILAC, BUTTER, SKY, PEACH = (
    (255, 214, 226), (205, 240, 222), (225, 214, 250), (255, 243, 196), (205, 230, 250), (255, 222, 200)
)
INK = (90, 60, 80)


def _font(size: int):
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
    ):
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _face(draw: ImageDraw.ImageDraw, cx: int, cy: int, scale: float = 1.0):
    """Two eyes, blush and a little smile: the kawaii bit."""
    e = int(9 * scale)
    gap = int(28 * scale)
    for dx in (-gap, gap):
        draw.ellipse((cx + dx - e, cy - e, cx + dx + e, cy + e), fill=INK)
        draw.ellipse((cx + dx - e // 3 + 3, cy - e // 2, cx + dx + e // 3 + 3, cy), fill="white")
        blush = int(12 * scale)
        by = cy + int(20 * scale)
        draw.ellipse((cx + dx * 1.6 - blush, by - blush // 2, cx + dx * 1.6 + blush, by + blush // 2), fill=(255, 150, 170))
    w = int(14 * scale)
    draw.arc((cx - w, cy + int(4 * scale), cx + w, cy + int(26 * scale)), 20, 160, fill=INK, width=max(3, int(4 * scale)))


def _sparkles(draw, colour):
    for x, y, r in ((50, 60, 10), (350, 70, 8), (330, 330, 11), (60, 320, 7)):
        draw.polygon([(x, y - r), (x + r // 3, y - r // 3), (x + r, y), (x + r // 3, y + r // 3),
                      (x, y + r), (x - r // 3, y + r // 3), (x - r, y), (x - r // 3, y - r // 3)], fill=colour)


def picture(kind: str, background, label: str) -> bytes:
    img = Image.new("RGB", (400, 400), background)
    d = ImageDraw.Draw(img)
    _sparkles(d, (255, 255, 255))
    if kind.startswith("notebook"):
        d.rounded_rectangle((110, 70, 290, 310), 18, fill=(255, 255, 255), outline=INK, width=5)
        for y in range(95, 300, 26):
            d.ellipse((98, y, 122, y + 12), outline=INK, width=4)
        if kind.endswith("dots"):
            for y in range(240, 300, 18):
                for x in range(140, 280, 18):
                    d.ellipse((x - 2, y - 2, x + 2, y + 2), fill=(200, 180, 200))
        else:
            for y in range(240, 300, 18):
                d.line((135, y, 275, y), fill=(200, 180, 200), width=3)
        _face(d, 200, 160)
    elif kind == "mini":
        d.rounded_rectangle((140, 110, 260, 290), 14, fill=(255, 255, 255), outline=INK, width=5)
        _face(d, 200, 190, 0.8)
    elif kind == "pen":
        for i, c in enumerate(((255, 130, 160), (130, 200, 170), (160, 140, 230), (250, 200, 90), (120, 180, 240))):
            x = 110 + i * 42
            d.rounded_rectangle((x, 90, x + 28, 290), 10, fill=c, outline=INK, width=3)
            d.polygon([(x, 290), (x + 28, 290), (x + 14, 320)], fill=(255, 255, 255), outline=INK)
        _face(d, 200, 180, 0.7)
    elif kind == "tape":
        for i, c in enumerate(((255, 160, 185), (150, 215, 190), (190, 170, 240))):
            cx, cy = 130 + i * 70, 200 + (i % 2) * 30
            d.ellipse((cx - 55, cy - 55, cx + 55, cy + 55), fill=c, outline=INK, width=4)
            d.ellipse((cx - 22, cy - 22, cx + 22, cy + 22), fill=background, outline=INK, width=3)
        _face(d, 200, 110, 0.6)
    elif kind == "stickers":
        d.rounded_rectangle((90, 70, 310, 330), 14, fill=(255, 255, 255), outline=INK, width=5)
        for i in range(6):
            x, y = 130 + (i % 3) * 70, 120 + (i // 3) * 110
            if i % 2:
                d.ellipse((x - 25, y - 25, x + 25, y + 25), fill=(255, 190, 205))
            else:
                pts = [(x + 28 * math.cos(math.radians(a)) * (1 if k % 2 == 0 else 0.45),
                        y + 28 * math.sin(math.radians(a)) * (1 if k % 2 == 0 else 0.45))
                       for k, a in enumerate(range(-90, 270, 36))]
                d.polygon(pts, fill=(250, 210, 100))
        _face(d, 200, 225, 0.6)
    elif kind == "book":
        d.polygon([(90, 110), (200, 90), (200, 310), (90, 330)], fill=(255, 255, 255), outline=INK)
        d.polygon([(200, 90), (310, 110), (310, 330), (200, 310)], fill=(255, 255, 255), outline=INK)
        d.line((200, 90, 200, 310), fill=INK, width=5)
        for y in range(140, 290, 30):
            d.line((110, y, 185, y - 6), fill=(200, 180, 200), width=4)
        _face(d, 255, 200, 0.6)
    elif kind == "case":
        d.rounded_rectangle((70, 150, 330, 300), 50, fill=(255, 255, 255), outline=INK, width=5)
        for ex in (140, 260):
            d.ellipse((ex - 25, 70, ex + 25, 170), fill=(255, 255, 255), outline=INK, width=5)
            d.ellipse((ex - 12, 90, ex + 12, 150), fill=(255, 190, 205))
        d.line((90, 190, 310, 190), fill=INK, width=4)
        _face(d, 200, 240, 0.8)
    elif kind == "eraser":
        d.polygon([(200, 90), (320, 290), (80, 290)], fill=(255, 255, 255), outline=INK)
        d.rectangle((150, 240, 250, 290), fill=(60, 60, 70))
        _face(d, 200, 200, 0.7)
    elif kind == "planner":
        d.rounded_rectangle((100, 70, 300, 320), 16, fill=(255, 255, 255), outline=INK, width=5)
        d.rectangle((100, 70, 300, 130), fill=(255, 170, 190), outline=INK, width=5)
        d.text((200, 100), "2027", font=_font(32), fill="white", anchor="mm")
        _face(d, 200, 215, 0.8)
    elif kind == "highlighter":
        for i, c in enumerate(((255, 180, 200), (190, 235, 200), (200, 190, 250), (255, 235, 150), (180, 215, 250), (255, 200, 170))):
            x = 95 + i * 36
            d.rounded_rectangle((x, 110, x + 30, 300), 8, fill=c, outline=INK, width=3)
        _face(d, 200, 80, 0.5)
    d.text((200, 370), label, font=_font(22), fill=INK, anchor="mm")
    buffer = io.BytesIO()
    img.save(buffer, "PNG", optimize=True)
    return buffer.getvalue()


# SKU, product, details, sold as, pack contents, price, picture (kind, colour) or None
PRODUCTS = [
    ("MK-NB-101", "Mochi Bunny Notebook", "A5, dotted, 120 pages, 100 gsm, pink soft cover", "notebook", "1", "1,250", ("notebook-dots", PINK)),
    ("MK-NB-102", "Mochi Bunny Notebook", "A5, lined, 120 pages, 100 gsm, mint soft cover", "notebook", "1", "1,250", ("notebook-lines", MINT)),
    ("MK-NB-110", "Kitty Cloud Mini Notebook", "A6, blank, 80 pages, lilac cover", "notebook", "1", "650", ("mini", LILAC)),
    ("MK-PL-201", "Sakura Days 2027 Planner", "A5, weekly and monthly pages, stickers inside, hardcover", "planner", "1", "3,400", ("planner", PINK)),
    ("MK-PN-301", "Pastel Dream Gel Pens", "0.5 mm, 10 pastel colours", "pack", "10", "1,100", ("pen", SKY)),
    ("MK-PN-305", "Boba Tea Highlighters", "chisel tip, 6 pastel colours", "set", "6", "950", ("highlighter", BUTTER)),
    ("MK-TP-401", "Strawberry Milk Washi Tape", "15 mm × 5 m per roll, 5 designs", "set", "5", "850", ("tape", PEACH)),
    ("MK-ST-501", "Lucky Star Sticker Sheets", "glossy vinyl, waterproof, 6 different sheets", "pack", "6", "600", ("stickers", LILAC)),
    ("MK-ER-601", "Onigiri Friends Erasers", "dust-free, 4 characters", "pack", "4", "420", ("eraser", MINT)),
    ("MK-PC-701", "Bunny Ears Pencil Case", "plush, zip, holds 30 pens", "unit", "1", "1,800", ("case", PINK)),
    ("MK-BK-801", "The Little Cat Café", "cosy novel, paperback, 240 pages, English", "book", "1", "1,950", ("book", PEACH)),
    ("MK-BK-802", "Doodle Every Day", "sketch-a-day activity book, 180 prompts, spiral bound", "book", "1", "1,450", ("book", SKY)),
    ("MK-BK-803", "Moonlight Bakery Vol. 1", "illustrated comic, paperback, 160 pages", "book", "1", "1,650", None),
    ("MK-PA-901", "Printer Paper A4", "80 gsm, white, 500 sheets", "ream", "500", "1,900", None),
    ("MK-GW-950", "Gift Wrapping", "pastel paper, ribbon and a handwritten tag", "item", "1", "150", None),
]

PARAGRAPHS = {
    "About us": [
        "Miku's Stationery is a small kawaii stationery and book shop in Karachi, Pakistan. "
        "Shop 7, Clifton Block 5, Karachi. We sell cute notebooks, pens, planners, stickers "
        "and cosy books, and we ship all over Pakistan.",
        "Opening hours: Monday to Saturday, 11:00 AM–8:00 PM PKT. Closed on Sundays.",
        "Prices are in Pakistani Rupees (PKR) and include sales tax.",
    ],
    "How we talk": [
        "We are cheerful, warm and a little bit kawaii. Short, friendly messages with a cute "
        "emoticon now and then, like (◕‿◕) or ✿. We never pressure anyone and we never promise "
        "anything we can't do.",
    ],
    "Delivery": [
        "Delivery within Karachi is PKR 250 for orders below PKR 3,000, and free for orders of "
        "PKR 3,000 or more.",
        "Delivery to the rest of Pakistan is PKR 350 for orders below PKR 5,000, and free for "
        "orders of PKR 5,000 or more.",
        "Karachi orders arrive in 1–2 working days; other cities in 3–5 working days. We do not "
        "offer same-day delivery.",
    ],
    "Discounts": [
        "School and office orders: orders of PKR 10,000 or more receive 10% off. There are no "
        "other discounts, and discount codes from other shops are not accepted.",
        "Gift wrapping is PKR 150 per item and is not discounted.",
    ],
    "Payment": [
        "Cash on Delivery is available for orders up to PKR 15,000.",
        "Orders above PKR 15,000 need 50% advance payment by bank transfer, JazzCash or Easypaisa "
        "before dispatch, and the rest on delivery.",
        "We do not offer instalments or pay-later.",
    ],
    "Returns and exchanges": [
        "Unused items in their original packaging can be exchanged within 7 days of delivery.",
        "Opened sticker sheets, opened washi tape and used notebooks cannot be returned.",
        "If a book arrives damaged, send us a photo within 48 hours of delivery and we will "
        "replace it.",
        "Refunds for approved returns are sent by bank transfer within 5 working days.",
    ],
    "Stock": [
        "The Sakura Days 2027 Planner is limited to 40 copies this season. Moonlight Bakery Vol. 2 "
        "is not out yet, and we do not take pre-orders.",
    ],
}


def build() -> Path:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MIKU'S STATIONERY ✿")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0xD9, 0x5F, 0x8A)
    sub = document.add_paragraph("Cute things for writing, planning and reading (◕‿◕)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "Test document — the shop, products and prices are made up. Price list 2026-10 | Currency: PKR"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    for heading in ("About us", "How we talk"):
        document.add_heading(heading, level=2)
        for text in PARAGRAPHS[heading]:
            document.add_paragraph(text)

    document.add_heading("Price list", level=2)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells,
        ("Photo", "SKU", "Product", "Details", "Sale Unit", "Pack", "Price (PKR)"),
    ):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for sku, name, details, unit, pack, price, art in PRODUCTS:
        cells = table.add_row().cells
        if art:
            kind, colour = art
            cells[0].paragraphs[0].add_run().add_picture(
                io.BytesIO(picture(kind, colour, " ".join([w for w in name.split() if w != "The"][:2]))),
                width=Cm(2.6),
            )
        for cell, text in zip(cells[1:], (sku, name, details, unit, pack, price)):
            cell.text = text

    for heading in ("Delivery", "Discounts", "Payment", "Returns and exchanges", "Stock"):
        document.add_heading(heading, level=2)
        for text in PARAGRAPHS[heading]:
            document.add_paragraph(text)

    document.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
