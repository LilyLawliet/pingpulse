"""Reading a client's own files, which is what actually unblocks onboarding.

Every shop has its catalogue written down already. Until this existed, taking
one on meant somebody retyping a PDF into .txt files — the real bottleneck in
signing a client, and nothing to do with the model.

What these protect, in order of how expensive the failure is:

  * a table row stays one line. The fact binding a product to its price *is*
    the row; flattened cell by cell the price survives and the binding does
    not, and the agent confidently quotes the wrong number. That reads as the
    model hallucinating when it was the parser;
  * a scanned PDF is refused with a reason, rather than indexed as nothing and
    leaving the client to discover the agent knows none of their products;
  * long documents are split before embedding, because one vector for forty
    pages points at the average of everything in it;
  * an upload lands in the uploader's tenant and nobody else's.
"""

from __future__ import annotations

import io

import pytest
from sqlalchemy import select

from app.models import KnowledgeDocument
from app.services import documents


# --------------------------------------------------------------- a real PDF
def _pdf(lines: list[str], pages: int = 1) -> bytes:
    """A minimal, valid PDF with a text layer.

    Written by hand rather than with a PDF library so the test suite does not
    grow a dependency purely to produce fixtures. It is the smallest thing
    pdfplumber will read as a page of text.
    """
    objects: list[bytes] = []
    page_ids = [4 + i * 2 for i in range(pages)]

    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(pages).encode() + b" >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for page in range(pages):
        stream_id = page_ids[page] + 1
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents "
            + str(stream_id).encode()
            + b" 0 R >>"
        )
        body = b"BT /F1 12 Tf 40 740 Td 14 TL\n"
        for line in lines:
            escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            body += b"(" + escaped.encode("latin-1", "replace") + b") Tj T*\n"
        body += b"ET"
        objects.append(b"<< /Length " + str(len(body)).encode() + b" >>\nstream\n" + body + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, payload in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + payload + b"\nendobj\n"

    xref_at = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_at).encode() + b"\n%%EOF\n"
    )
    return bytes(out)


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ extraction
def test_a_pdf_price_list_is_read():
    data = _pdf(["Aurora SoundPods Pro  $89  in stock 42", "Halo Desk Lamp  $59  in stock 18"])

    result = documents.extract("prices.pdf", data)

    assert result.kind == "pdf"
    assert result.pages == 1
    assert "SoundPods Pro" in result.text
    assert "$89" in result.text


def test_a_word_document_is_read():
    data = _docx(["Delivery is next day in the city.", "Returns within 14 days."])

    result = documents.extract("policies.docx", data)

    assert result.kind == "docx"
    assert "next day" in result.text
    assert "14 days" in result.text


def test_a_table_keeps_each_row_on_one_line():
    """The expensive one.

    A price belongs to the product it shares a row with. Emitted cell by cell
    both survive and the binding between them does not, and the agent quotes a
    real price for the wrong product — which looks like the model inventing
    things when the parser threw the relationship away.
    """
    data = _docx(
        ["Catalogue"],
        table=[["Item", "Price"], ["SoundPods Pro", "$89"], ["Trail Backpack", "$69"]],
    )

    text = documents.extract("catalogue.docx", data).text

    assert "SoundPods Pro | $89" in text
    assert "Trail Backpack | $69" in text
    # And never the other way round.
    assert "SoundPods Pro | $69" not in text


def test_a_scanned_pdf_is_refused_with_a_reason():
    """Photographed pages have no text layer. Indexing that quietly gives the
    client an agent that knows nothing about their products and no clue why."""
    blank = _pdf([""], pages=3)

    with pytest.raises(documents.UnreadableDocument) as raised:
        documents.extract("scan.pdf", blank)

    assert "scan" in str(raised.value).lower()


def test_an_unsupported_file_says_what_is_supported():
    with pytest.raises(documents.UnsupportedDocument) as raised:
        documents.extract("catalogue.xlsx", b"whatever")

    assert ".pdf" in str(raised.value)


def test_a_corrupt_file_is_a_readable_error_not_a_stack_trace():
    with pytest.raises(documents.UnreadableDocument):
        documents.extract("broken.pdf", b"not a pdf at all")


# ------------------------------------------------------------------- chunking
def test_a_long_document_is_split_for_embedding():
    """One vector for a whole document points at the average of everything."""
    text = "\n\n".join(f"Paragraph {i}. " + ("filler words here. " * 20) for i in range(12))

    passages = documents.chunk(text)

    assert len(passages) > 1
    assert all(len(p) <= documents.TARGET_CHARS * 1.5 for p in passages)
    assert "Paragraph 0." in passages[0]
    assert "Paragraph 11." in passages[-1]


def test_a_table_is_never_split_down_the_middle():
    """Half a price list is worse than two of them."""
    rows = "\n".join(f"Product {i} | ${i * 7} | in stock" for i in range(90))

    passages = documents.chunk(f"Catalogue\n\n{rows}")

    holding = [p for p in passages if "Product 0 |" in p]
    assert holding, "the table vanished"
    assert "Product 89 |" in holding[0], "the table was cut in half"


def test_a_short_document_stays_one_passage():
    assert len(documents.chunk("Delivery is next day. Returns within 14 days.")) == 1


def test_titles_are_only_numbered_when_there_is_more_than_one():
    assert documents.title_for("price-list.pdf", 0, 1) == "Price List"
    assert documents.title_for("price-list.pdf", 0, 3) == "Price List (1/3)"


# ------------------------------------------------------------------ the upload
@pytest.mark.asyncio
async def test_uploading_a_catalogue_makes_it_searchable(org_a, db_session):
    """End to end: the thing that replaces retyping a client's price list."""
    data = _docx(
        ["Aurora Retail catalogue"],
        table=[["Item", "Price"], ["SoundPods Pro", "$89"], ["Halo Desk Lamp", "$59"]],
    )

    response = await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("catalogue.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"doc_type": "product"},
    )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["passages_indexed"] >= 1
    assert body["tables_found"] == 1

    hits = await org_a.get("/api/v1/knowledge/search?q=how much are the soundpods")
    assert hits.status_code == 200
    assert any("89" in hit["content"] for hit in hits.json()), (
        "the price was indexed but is not retrievable"
    )


@pytest.mark.asyncio
async def test_an_upload_belongs_to_the_tenant_that_sent_it(org_a, org_b, db_session):
    """Knowledge is a tenant's own. Another shop's prices must never be quoted."""
    data = _docx(["Secret pricing: SoundPods Pro at cost is $41."])

    uploaded = await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("costs.docx", data, "application/octet-stream")},
    )
    assert uploaded.status_code == 201

    theirs = await org_b.get("/api/v1/knowledge/search?q=cost price soundpods")
    assert theirs.json() == [], "another tenant could read this shop's cost prices"

    stored = (
        await db_session.execute(
            select(KnowledgeDocument).where(
                KnowledgeDocument.organization_id == org_b.organization_id
            )
        )
    ).scalars().all()
    assert stored == []


@pytest.mark.asyncio
async def test_a_scanned_upload_is_rejected_in_plain_words(org_a):
    """The shop owner reading this is not an engineer."""
    response = await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("scan.pdf", _pdf([""], pages=3), "application/pdf")},
    )

    assert response.status_code == 422
    assert "scan" in response.json()["detail"].lower()


@pytest.mark.asyncio
async def test_an_empty_file_is_refused(org_a):
    response = await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert response.status_code == 422


# ------------------------------------------------- what the dashboard shows
@pytest.mark.asyncio
async def test_uploads_are_listed_by_file_not_by_passage(org_a):
    """Someone who uploaded one price list should see one row, not twelve.

    Splitting into passages is right for retrieval and meaningless to the
    person who sent the file.
    """
    long_text = "\n\n".join(f"Section {i}. " + ("detail " * 60) for i in range(8))

    uploaded = await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("handbook.txt", long_text.encode(), "text/plain")},
    )
    assert uploaded.status_code == 201
    assert uploaded.json()["passages_indexed"] > 1, "not enough text to prove the point"

    listed = await org_a.get("/api/v1/knowledge/sources")
    assert listed.status_code == 200
    rows = listed.json()
    assert len(rows) == 1, "one file should be one row"
    assert rows[0]["source"] == "handbook.txt"
    assert rows[0]["passages"] == uploaded.json()["passages_indexed"]


@pytest.mark.asyncio
async def test_removing_a_file_removes_all_of_it(org_a):
    """Half a price list is worse than none: the agent quotes the half that
    survived and looks confidently wrong."""
    data = _docx(["Winter prices", "Scarf $34", "Boots $189"])
    await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("winter.docx", data, "application/octet-stream")},
    )

    removed = await org_a._client.delete(
        "/api/v1/knowledge/sources?source=winter.docx", headers=org_a.headers
    )
    assert removed.status_code == 204

    assert (await org_a.get("/api/v1/knowledge/sources")).json() == []
    assert (await org_a.get("/api/v1/knowledge/search?q=scarf")).json() == []


@pytest.mark.asyncio
async def test_one_tenant_cannot_delete_anothers_upload(org_a, org_b):
    """The filename is the handle, and filenames collide between shops."""
    await org_a._client.post(
        "/api/v1/knowledge/upload",
        headers=org_a.headers,
        files={"file": ("prices.txt", b"Scarf $34 in stock", "text/plain")},
    )

    theirs = await org_b._client.delete(
        "/api/v1/knowledge/sources?source=prices.txt", headers=org_b.headers
    )
    assert theirs.status_code == 404, "a shared filename must not delete across tenants"

    still_there = await org_a.get("/api/v1/knowledge/sources")
    assert len(still_there.json()) == 1


# ------------------------------------------------- markdown is for the eyes
def test_markdown_syntax_does_not_reach_the_customer():
    """What is stored here is quoted a sentence at a time, asterisks and all.

    A client's knowledge base had agent scripting in it and the agent read it
    back word for word - "OPEN: 'Absolutely — Constrivo Group handles...'".
    Markdown is the same fault in a different form: "**Kitchen remodeling.**
    Custom designs" is a heading and a bold phrase to a reader, and two
    asterisks and a full stop to a customer.
    """
    from app.services.documents import extract

    source = (
        "# Constrivo Group\n\n"
        "## Services\n\n"
        "**Kitchen remodeling.** Custom designs, premium materials and expert installation.\n\n"
        "- Home additions\n"
        "- Roofing\n\n"
        "> Build Now. Pay Over Time.\n\n"
        "| Service | Price |\n|---|---|\n| Kitchen | free estimate |\n\n"
        "See [our site](https://constrivogroup.com/) for more.\n\n"
        "---\n"
    )
    text = extract("knowledge-base.md", source.encode("utf-8")).text

    for syntax in ("#", "**", "](", "|", "> "):
        assert syntax not in text, f"{syntax!r} survived into the stored text"
    # The words are all still there.
    assert "Kitchen remodeling. Custom designs" in text
    assert "Home additions" in text
    assert "Build Now. Pay Over Time." in text
    assert "our site" in text
    assert "free estimate" in text


def test_a_plain_text_file_is_left_alone():
    from app.services.documents import extract

    plain = "Constrivo Group works in Miami and South Florida.\nCall +1 305-748-3629."
    assert extract("notes.txt", plain.encode("utf-8")).text == plain
